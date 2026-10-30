"""In-memory extraction and layout-preserving redaction per file format.

Every handler exposes:
    extract(data) -> list[Segment]
    redact(data, plan) -> bytes   where plan = {segment_id: [Span, ...]}
Segment ids are produced by a deterministic traversal, so redact() can find the same text node again.
"""
import email
import io
from dataclasses import dataclass
from email import policy
from pathlib import Path

import pymupdf as fitz
from docx import Document as DocxDocument
from docx.text.run import Run
from openpyxl import load_workbook
from pptx import Presentation

from .models import Segment


@dataclass
class Span:
    start: int
    end: int
    replacement: str
    original: str
    entity_type: str = ""


SUPPORTED = {".txt", ".csv", ".md", ".json", ".log", ".docx", ".xlsx", ".pptx", ".pdf", ".eml"}


def ext_of(name: str) -> str:
    return Path(name).suffix.lower()


def is_supported(name: str) -> bool:
    return ext_of(name) in SUPPORTED


def extract(name: str, data: bytes) -> list[Segment]:
    return _HANDLERS[_kind(name)][0](data)


def redact(name: str, data: bytes, plan: dict[str, list[Span]]) -> bytes:
    return _HANDLERS[_kind(name)][1](data, plan)


def _kind(name: str) -> str:
    ext = ext_of(name)
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {ext}")
    return "text" if ext in {".txt", ".csv", ".md", ".json", ".log"} else ext[1:]


def _apply_spans(text: str, spans: list[Span]) -> str:
    for s in sorted(spans, key=lambda s: s.start, reverse=True):
        text = text[: s.start] + s.replacement + text[s.end:]
    return text


def _apply_to_runs(runs, spans: list[Span]) -> None:
    """Redact spans that may cross run boundaries while keeping each run's formatting.

    The replacement is written into the first run the span touches; the remaining
    characters are removed from subsequent runs.
    """
    texts = [r.text or "" for r in runs]
    bounds, pos = [], 0
    for t in texts:
        bounds.append((pos, pos + len(t)))
        pos += len(t)
    new = list(texts)
    for s in sorted(spans, key=lambda s: s.start, reverse=True):
        first = True
        for i, (rs, re_) in enumerate(bounds):
            if re_ <= s.start or rs >= s.end:
                continue
            a, b = max(s.start, rs) - rs, min(s.end, re_) - rs
            new[i] = new[i][:a] + (s.replacement if first else "") + new[i][b:]
            first = False
    for run, old, value in zip(runs, texts, new):
        if old != value:
            run.text = value


# ---------- plain text ----------
def _decode(data: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "latin-1"


def _text_extract(data: bytes) -> list[Segment]:
    text, _ = _decode(data)
    return [Segment("body", text, "Document body")]


def _text_redact(data: bytes, plan: dict[str, list[Span]]) -> bytes:
    text, enc = _decode(data)
    enc = "utf-8" if enc == "utf-8-sig" else enc
    return _apply_spans(text, plan.get("body", [])).encode(enc, errors="replace")


# ---------- Word ----------
def _docx_paragraphs(doc):
    """Yield (segment_id, label, paragraph) in a stable order, de-duplicating merged table cells."""
    seen_cells = set()

    def walk_tables(tables, prefix):
        for ti, table in enumerate(tables):
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    if id(cell._tc) in seen_cells:
                        continue
                    seen_cells.add(id(cell._tc))
                    for pi, p in enumerate(cell.paragraphs):
                        yield f"{prefix}t{ti}r{ri}c{ci}p{pi}", f"Table {ti + 1}, row {ri + 1}", p
                    yield from walk_tables(cell.tables, f"{prefix}t{ti}r{ri}c{ci}.")

    for i, p in enumerate(doc.paragraphs):
        yield f"p{i}", f"Paragraph {i + 1}", p
    yield from walk_tables(doc.tables, "")
    for si, section in enumerate(doc.sections):
        for kind, part in (("hdr", section.header), ("ftr", section.footer)):
            if part.is_linked_to_previous:
                continue
            for pi, p in enumerate(part.paragraphs):
                yield f"s{si}{kind}p{pi}", f"Section {si + 1} {'header' if kind == 'hdr' else 'footer'}", p
            yield from walk_tables(part.tables, f"s{si}{kind}.")


def _docx_runs(paragraph) -> list[Run]:
    # Includes runs nested in hyperlinks / tracked insertions, which paragraph.runs skips.
    return [Run(r, paragraph) for r in paragraph._p.xpath(".//w:r[not(ancestor::w:del)]")]


def _docx_extract(data: bytes) -> list[Segment]:
    doc = DocxDocument(io.BytesIO(data))
    out = []
    for sid, label, p in _docx_paragraphs(doc):
        text = "".join(r.text or "" for r in _docx_runs(p))
        if text.strip():
            out.append(Segment(sid, text, label))
    return out


def _docx_redact(data: bytes, plan: dict[str, list[Span]]) -> bytes:
    doc = DocxDocument(io.BytesIO(data))
    for sid, _, p in _docx_paragraphs(doc):
        if sid in plan:
            _apply_to_runs(_docx_runs(p), plan[sid])
    originals = [s.original for spans in plan.values() for s in spans]
    for rel in doc.part.rels.values():
        if rel.is_external and any(o in rel.target_ref for o in originals):
            rel._target = "#redacted"
    _scrub_core_props(doc.core_properties)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _scrub_core_props(props) -> None:
    for attr in ("author", "last_modified_by", "comments", "keywords", "subject"):
        try:
            setattr(props, attr, "")
        except Exception:
            pass


# ---------- Excel ----------
def _xlsx_extract(data: bytes) -> list[Segment]:
    wb = load_workbook(io.BytesIO(data))
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.strip() and not v.startswith("="):
                    out.append(Segment(f"{ws.title}!{cell.coordinate}", v, f"{ws.title}!{cell.coordinate}"))
                elif isinstance(v, int) and len(str(v)) >= 9:  # long numbers: phone / Aadhaar / account no.
                    out.append(Segment(f"{ws.title}!{cell.coordinate}", str(v), f"{ws.title}!{cell.coordinate}"))
    return out


def _xlsx_redact(data: bytes, plan: dict[str, list[Span]]) -> bytes:
    wb = load_workbook(io.BytesIO(data))
    for sid, spans in plan.items():
        sheet, coord = sid.rsplit("!", 1)
        cell = wb[sheet][coord]
        cell.value = _apply_spans(str(cell.value), spans)
    _scrub_core_props(wb.properties)
    for attr in ("creator", "lastModifiedBy"):
        setattr(wb.properties, attr, None)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- PowerPoint ----------
def _pptx_paragraphs(prs):
    def walk_shapes(shapes, prefix):
        for shi, shape in enumerate(shapes):
            sp = f"{prefix}sh{shi}"
            if shape.shape_type == 6:  # group
                yield from walk_shapes(shape.shapes, sp + ".")
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for pi, p in enumerate(shape.text_frame.paragraphs):
                    yield f"{sp}p{pi}", p
            if getattr(shape, "has_table", False) and shape.has_table:
                for ri, row in enumerate(shape.table.rows):
                    for ci, cell in enumerate(row.cells):
                        for pi, p in enumerate(cell.text_frame.paragraphs):
                            yield f"{sp}r{ri}c{ci}p{pi}", p

    for si, slide in enumerate(prs.slides):
        for sid, p in walk_shapes(slide.shapes, f"s{si}."):
            yield sid, f"Slide {si + 1}", p
        if slide.has_notes_slide:
            for pi, p in enumerate(slide.notes_slide.notes_text_frame.paragraphs):
                yield f"s{si}.notes.p{pi}", f"Slide {si + 1} notes", p


def _pptx_extract(data: bytes) -> list[Segment]:
    prs = Presentation(io.BytesIO(data))
    out = []
    for sid, label, p in _pptx_paragraphs(prs):
        text = "".join(r.text for r in p.runs)
        if text.strip():
            out.append(Segment(sid, text, label))
    return out


def _pptx_redact(data: bytes, plan: dict[str, list[Span]]) -> bytes:
    prs = Presentation(io.BytesIO(data))
    for sid, _, p in _pptx_paragraphs(prs):
        if sid in plan:
            _apply_to_runs(list(p.runs), plan[sid])
    _scrub_core_props(prs.core_properties)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------- PDF ----------
def _pdf_extract(data: bytes) -> list[Segment]:
    with fitz.open(stream=data, filetype="pdf") as doc:
        return [Segment(f"page{i}", page.get_text(), f"Page {i + 1}")
                for i, page in enumerate(doc) if page.get_text().strip()]


def _pdf_redact(data: bytes, plan: dict[str, list[Span]]) -> bytes:
    """True redaction: underlying text is removed, a black box with a readable label is drawn."""
    with fitz.open(stream=data, filetype="pdf") as doc:
        for sid, spans in plan.items():
            page = doc[int(sid.removeprefix("page"))]
            for s in spans:
                needle = " ".join(s.original.split())
                label = s.replacement if len(s.replacement) <= 40 and "█" not in s.replacement else ""
                for rect in page.search_for(needle):
                    # Label must fit inside the box, otherwise neighbouring characters get removed.
                    size = min(7.0, rect.width / max(len(label), 1) / 0.6, rect.height * 0.7)
                    page.add_redact_annot(rect, text=label if size >= 3.5 else "", fill=(0, 0, 0),
                                          text_color=(1, 1, 1), fontsize=size if size >= 3.5 else 7)
            page.apply_redactions()
        doc.set_metadata({})
        return doc.tobytes(garbage=3, deflate=True)


# ---------- Email (.eml) ----------
_HEADERS = ("From", "To", "Cc", "Reply-To", "Subject")


def _eml_parts(msg):
    for i, part in enumerate(msg.walk()):
        if part.get_content_maintype() == "text" and part.get_content_disposition() != "attachment":
            yield f"part{i}", part


def _eml_extract(data: bytes) -> list[Segment]:
    msg = email.message_from_bytes(data, policy=policy.default)
    out = [Segment(f"hdr:{h}", str(msg[h]), f"Header {h}") for h in _HEADERS if msg[h]]
    for sid, part in _eml_parts(msg):
        out.append(Segment(sid, part.get_content(), f"Body ({part.get_content_type()})"))
    return out


def _eml_redact(data: bytes, plan: dict[str, list[Span]]) -> bytes:
    msg = email.message_from_bytes(data, policy=policy.default)
    for h in _HEADERS:
        sid = f"hdr:{h}"
        if sid in plan and msg[h]:
            new_value = _apply_spans(str(msg[h]), plan[sid])
            del msg[h]
            msg[h] = new_value
    for sid, part in list(_eml_parts(msg)):
        if sid in plan:
            part.set_content(_apply_spans(part.get_content(), plan[sid]),
                             subtype=part.get_content_subtype(), charset="utf-8")
    return msg.as_bytes(policy=policy.SMTP)


_HANDLERS = {
    "text": (_text_extract, _text_redact),
    "docx": (_docx_extract, _docx_redact),
    "xlsx": (_xlsx_extract, _xlsx_redact),
    "pptx": (_pptx_extract, _pptx_redact),
    "pdf": (_pdf_extract, _pdf_redact),
    "eml": (_eml_extract, _eml_redact),
}
